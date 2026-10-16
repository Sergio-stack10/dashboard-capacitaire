"""Génère data/data.json depuis la roadmap (.xlsm) et le fichier Calcul CAPA (.xlsx).
Usage : python scripts/build_data.py RM_Auchan.xlsm Calcul_CAPA.xlsx"""
import json, sys, datetime
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI
RM, CP = sys.argv[1], sys.argv[2]
ACTS=[("Hyper","Voix",126,213,300),("C2E","Voix",127,214,301),("AFR","Voix",128,215,302),("Drive","Voix",130,217,304),
("SAV","Voix",131,218,305),("Hyper","Mail",139,226,313),("AFR","Mail",140,227,314),("Drive","Mail",142,229,316),
("Chat","Autres",149,236,323),("Borne/LSA","Autres",150,237,324)]
ws=load_workbook(RM,data_only=True)["AUCHAN"]
c0=next((c for c in range(1,200) if ws.cell(3,c).value=="S1__2026"),CI("BU"))
S=lambda r:[round(v,1) if isinstance(v:=ws.cell(r,c0+i).value,(int,float)) else 0 for i in range(52)]
roadmap=[dict(a=a,c=c,prev=S(p),etp=S(e),capa=S(k)) for a,c,p,e,k in ACTS]
dv=load_workbook(CP,data_only=True)["DAILY VISION"]
cols=[(c,dv.cell(15,c).value.date().isoformat()) for c in range(5,dv.max_column+1) if isinstance(dv.cell(15,c).value,datetime.datetime)]
cap=lambda s:s if len(s)<=3 or s=="Hyper" else s[0]+s[1:].lower()
fl=[]
for r in range(1,dv.max_row+1):
    if dv.cell(r,4).value=="Capacitaire en actes":
        c,a=cap(str(dv.cell(r,1).value).title() if str(dv.cell(r,1).value).isupper() else str(dv.cell(r,1).value)),cap(str(dv.cell(r,2).value))
        g=lambda o:[round(dv.cell(r+o,k).value or 0) if isinstance(dv.cell(r+o,k).value,(int,float)) else 0 for k,_ in cols]
        fl.append(dict(k=c+" "+a,c=c,a=a,prev=g(-2),contr=g(-1),capa=g(0)))
fl=[f for f in fl if any(f["prev"]) or any(f["capa"])]
nz=[i for i in range(len(cols)) if any(f["prev"][i] or f["capa"][i] for f in fl)]
lo,hi=nz[0],nz[-1]+1
capa=dict(dates=[d for _,d in cols][lo:hi],flux=[{**f,"prev":f["prev"][lo:hi],"contr":f["contr"][lo:hi],"capa":f["capa"][lo:hi]} for f in fl])
json.dump(dict(meta=dict(roadmap=RM.split("/")[-1],capa=CP.split("/")[-1],maj=datetime.date.today().isoformat()),roadmap=roadmap,capa=capa),open("data/data.json","w"),separators=(",",":"))
print("OK",len(roadmap),"activités;",len(fl),"flux;",len(capa["dates"]),"jours",capa["dates"][0],"->",capa["dates"][-1])
