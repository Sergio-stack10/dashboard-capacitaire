"""Extrait la roadmap (onglet AUCHAN) vers data/roadmap.json.
Usage : python scripts/extract_roadmap.py chemin/RM_Auchan.xlsm
Lignes lues : prévisions finales, besoin en ETP, capa en actes (52 semaines 2026, à partir de la colonne BU)."""
import json, sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string as CI

SRC = sys.argv[1] if len(sys.argv) > 1 else "RM_Auchan.xlsm"
OUT = "data/roadmap.json"
N, C0 = 52, CI("BU")
# (activité, channel, ligne prévision finale, ligne besoin ETP, ligne capa en actes)
ACTS = [("Hyper","Voix",126,213,300),("C2E","Voix",127,214,301),("AFR","Voix",128,215,302),
        ("Drive","Voix",130,217,304),("SAV","Voix",131,218,305),("Hyper","Mail",139,226,313),
        ("AFR","Mail",140,227,314),("Drive","Mail",142,229,316),
        ("Chat","Autres",149,236,323),("Borne/LSA","Autres",150,237,324)]

ws = load_workbook(SRC, data_only=True)["AUCHAN"]
def serie(r):
    return [round(v, 1) if isinstance(v := ws.cell(r, C0 + i).value, (int, float)) else 0 for i in range(N)]

data = [dict(a=a, c=c, prev=serie(p), etp=serie(e), capa=serie(k)) for a, c, p, e, k in ACTS]
json.dump(data, open(OUT, "w"), separators=(",", ":"))
print("OK ->", OUT, len(data), "activités")
